"""Unit tests for privacyfence.audit_log — the append-only decision trail."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from freezegun import freeze_time

from privacyfence import paths
from privacyfence.audit_log import (
    APPROVED_LIKE_DECISIONS,
    CURRENT_SCHEMA_VERSION,
    GENESIS_HASH,
    AuditEntry,
    AuditLogger,
    _previous_week,
    compute_security_config_hash,
    current_week,
    get_audit_logger,
    init_audit_logger,
)


def make_entry(**overrides) -> AuditEntry:
    defaults = dict(
        timestamp="2026-07-06T12:00:00+00:00",
        week="2026-W28",
        request_id="",
        connector="gmail",
        tool="gmail_get_message",
        tool_name="Read Gmail message",
        summary="Message from alice@example.com",
        sender="alice@example.com",
        decision="approved",
        auto_accept_rule="",
        latency_seconds=1.23,
    )
    defaults.update(overrides)
    return AuditEntry(**defaults)


class TestPiiDetectedField:
    def test_defaults_to_false(self):
        assert make_entry().pii_detected is False

    def test_round_trips_through_jsonl(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(pii_detected=True))

        line = (tmp_path / "2026-W28.jsonl").read_text(encoding="utf-8").splitlines()[0]
        assert json.loads(line)["pii_detected"] is True

    def test_old_jsonl_lines_without_the_field_still_parse(self):
        # Entries written before this field existed have no "pii_detected"
        # key at all; export_week_to_excel reconstructs AuditEntry(**line),
        # so the field needs a default rather than being required.
        legacy = dict(
            timestamp="2026-07-06T12:00:00+00:00", week="2026-W28", request_id="",
            connector="gmail", tool="gmail_get_message", tool_name="Read Gmail message",
            summary="s", sender="a@example.com", decision="approved",
            auto_accept_rule="", latency_seconds=1.0,
        )
        entry = AuditEntry(**legacy)
        assert entry.pii_detected is False


class TestPiiCategoriesField:
    def test_defaults_to_empty_list(self):
        assert make_entry().pii_categories == []

    def test_round_trips_through_jsonl(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(pii_categories=["IBAN (bank account number)", "IP address"]))

        line = (tmp_path / "2026-W28.jsonl").read_text(encoding="utf-8").splitlines()[0]
        assert json.loads(line)["pii_categories"] == ["IBAN (bank account number)", "IP address"]

    def test_old_jsonl_lines_without_the_field_still_parse(self):
        legacy = dict(
            timestamp="2026-07-06T12:00:00+00:00", week="2026-W28", request_id="",
            connector="gmail", tool="gmail_get_message", tool_name="Read Gmail message",
            summary="s", sender="a@example.com", decision="approved",
            auto_accept_rule="", latency_seconds=1.0,
        )
        entry = AuditEntry(**legacy)
        assert entry.pii_categories == []

    def test_each_entry_gets_its_own_list_not_a_shared_default(self):
        # dataclass mutable-default pitfall: default_factory=list must give
        # each AuditEntry its own list, not one shared across every instance
        # missing the kwarg.
        a = make_entry()
        b = make_entry()
        a.pii_categories.append("IBAN (bank account number)")
        assert b.pii_categories == []


class TestPiiMatchDetailsField:
    def test_defaults_to_empty_string(self):
        assert make_entry().pii_match_details == ""

    def test_round_trips_through_jsonl(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(pii_match_details="Salary/compensation information: salary"))

        line = (tmp_path / "2026-W28.jsonl").read_text(encoding="utf-8").splitlines()[0]
        assert json.loads(line)["pii_match_details"] == "Salary/compensation information: salary"

    def test_old_jsonl_lines_without_the_field_still_parse(self):
        legacy = dict(
            timestamp="2026-07-06T12:00:00+00:00", week="2026-W28", request_id="",
            connector="gmail", tool="gmail_get_message", tool_name="Read Gmail message",
            summary="s", sender="a@example.com", decision="approved",
            auto_accept_rule="", latency_seconds=1.0,
        )
        entry = AuditEntry(**legacy)
        assert entry.pii_match_details == ""


class TestApprovedLikeDecisionsIsPublic:
    def test_expected_members(self):
        assert APPROVED_LIKE_DECISIONS == {
            "approved", "auto_accepted", "accepted_via_accept_all", "accepted_via_temp_session",
        }


class TestClaudeReasonField:
    def test_defaults_to_empty_string(self):
        assert make_entry().claude_reason == ""

    def test_round_trips_through_jsonl(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(claude_reason="Summarizing the Q3 budget for the user."))

        line = (tmp_path / "2026-W28.jsonl").read_text(encoding="utf-8").splitlines()[0]
        assert json.loads(line)["claude_reason"] == "Summarizing the Q3 budget for the user."

    def test_old_jsonl_lines_without_the_field_still_parse(self):
        # Same backward-compatibility need as pii_detected above -- entries
        # written before this field existed have no "claude_reason" key.
        legacy = dict(
            timestamp="2026-07-06T12:00:00+00:00", week="2026-W28", request_id="",
            connector="gmail", tool="gmail_get_message", tool_name="Read Gmail message",
            summary="s", sender="a@example.com", decision="approved",
            auto_accept_rule="", latency_seconds=1.0,
        )
        entry = AuditEntry(**legacy)
        assert entry.claude_reason == ""


class TestCurrentWeek:
    @freeze_time("2026-07-06")  # a Monday, ISO week 28 of 2026
    def test_format(self):
        assert current_week() == "2026-W28"

    @freeze_time("2026-01-01")  # ISO week boundary: this date is ISO week 1 of 2026
    def test_iso_week_boundary(self):
        assert current_week() == "2026-W01"


class TestAuditLoggerRecord:
    def test_record_appends_jsonl_line(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry())

        week_file = tmp_path / "2026-W28.jsonl"
        assert week_file.exists()
        lines = week_file.read_text(encoding="utf-8").splitlines()
        assert len(lines) == 1
        data = json.loads(lines[0])
        assert data["connector"] == "gmail"
        assert data["decision"] == "approved"

    def test_record_appends_multiple_entries_to_same_week(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(decision="approved"))
        logger.record(make_entry(decision="rejected"))

        lines = (tmp_path / "2026-W28.jsonl").read_text(encoding="utf-8").splitlines()
        assert len(lines) == 2
        assert json.loads(lines[0])["decision"] == "approved"
        assert json.loads(lines[1])["decision"] == "rejected"

    def test_record_separates_different_weeks(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(week="2026-W01"))
        logger.record(make_entry(week="2026-W28"))

        assert (tmp_path / "2026-W01.jsonl").exists()
        assert (tmp_path / "2026-W28.jsonl").exists()

    def test_log_dir_created_if_missing(self, tmp_path):
        target = tmp_path / "nested" / "audit"
        AuditLogger(str(target))
        assert target.is_dir()


class TestExportWeekToExcel:
    def test_export_returns_none_when_no_jsonl_exists(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        assert logger.export_week_to_excel("2026-W28") is None

    def test_export_produces_workbook_with_expected_content(self, tmp_path):
        openpyxl = pytest.importorskip("openpyxl")

        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(
            decision="approved", pii_detected=True,
            pii_categories=["IBAN (bank account number)"],
            pii_match_details="IBAN (bank account number): DE••••••••••••••••••00",
            claude_reason="Summarizing for the user.",
        ))
        logger.record(make_entry(decision="auto_accepted", auto_accept_rule="i_am_sender"))
        logger.record(make_entry(decision="rejected"))

        output = logger.export_week_to_excel("2026-W28")
        assert output == str(tmp_path / "2026-W28.xlsx")

        wb = openpyxl.load_workbook(output)
        ws = wb["Decisions"]
        assert ws.cell(row=1, column=1).value == "Timestamp"
        # 3 data rows + 1 header row
        assert ws.max_row == 4
        decisions_col = [ws.cell(row=r, column=8).value for r in range(2, 5)]
        assert decisions_col == ["approved", "auto_accepted", "rejected"]

        assert ws.cell(row=1, column=11).value == "PII Detected"
        pii_col = [ws.cell(row=r, column=11).value for r in range(2, 5)]
        assert pii_col == ["Yes", None, None]  # openpyxl reads back "" cells as None

        assert ws.cell(row=1, column=12).value == "PII Categories"
        categories_col = [ws.cell(row=r, column=12).value for r in range(2, 5)]
        assert categories_col == ["IBAN (bank account number)", None, None]

        assert ws.cell(row=1, column=13).value == "PII Match Details"
        details_col = [ws.cell(row=r, column=13).value for r in range(2, 5)]
        assert details_col == ["IBAN (bank account number): DE••••••••••••••••••00", None, None]

        assert ws.cell(row=1, column=14).value == "Claude's Reason (unverified)"
        reason_col = [ws.cell(row=r, column=14).value for r in range(2, 5)]
        assert reason_col == ["Summarizing for the user.", None, None]

        summary = wb["Summary"]
        summary_rows = {row[0].value: row[1].value for row in summary.iter_rows(min_row=2) if row[0].value}
        assert summary_rows["Total decisions"] == 3
        assert summary_rows["Approved (manual)"] == 1
        assert summary_rows["Auto-accepted"] == 1
        assert summary_rows["Rejected"] == 1
        assert summary_rows["PII flagged (any decision)"] == 1
        assert summary_rows["IBAN (bank account number)"] == 1

    def test_export_omits_category_breakdown_when_no_entry_has_categories(self, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl

        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(decision="approved"))

        output = logger.export_week_to_excel("2026-W28")
        wb = openpyxl.load_workbook(output)
        summary = wb["Summary"]
        labels = [row[0].value for row in summary.iter_rows(min_row=2) if row[0].value]
        assert "By PII category (refinement trial)" not in labels

    def test_export_skips_malformed_lines(self, tmp_path):
        pytest.importorskip("openpyxl")
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry())
        week_file = tmp_path / "2026-W28.jsonl"
        with open(week_file, "a", encoding="utf-8") as fh:
            fh.write("not valid json\n")
            fh.write("\n")  # blank line

        output = logger.export_week_to_excel("2026-W28")
        assert output is not None

    def test_export_empty_file_returns_none(self, tmp_path):
        pytest.importorskip("openpyxl")
        logger = AuditLogger(str(tmp_path))
        week_file = tmp_path / "2026-W28.jsonl"
        week_file.write_text("", encoding="utf-8")
        assert logger.export_week_to_excel("2026-W28") is None


class TestRecentEntries:
    @freeze_time("2026-07-06")  # ISO week 28 of 2026
    def test_skips_blank_lines(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(week=current_week()))
        with open(tmp_path / f"{current_week()}.jsonl", "a", encoding="utf-8") as fh:
            fh.write("\n")

        assert len(logger.recent_entries()) == 1

    @freeze_time("2026-07-06")
    def test_skips_malformed_lines(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(week=current_week()))
        with open(tmp_path / f"{current_week()}.jsonl", "a", encoding="utf-8") as fh:
            fh.write("not valid json\n")

        assert len(logger.recent_entries()) == 1

    @freeze_time("2026-07-06")
    def test_stops_at_current_week_once_limit_reached(self, tmp_path):
        # Never falls through to the previous week's file once the current
        # week alone already satisfies `limit`.
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(week=current_week()))
        logger.record(make_entry(week=_previous_week(current_week())))

        assert len(logger.recent_entries(limit=1)) == 1


class TestRecentMatches:
    """The request-fingerprint feature: (connector, tool, summary) counted
    against one week's approved-like decisions."""

    def test_counts_matching_approved_entries(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        for _ in range(3):
            logger.record(make_entry(week="2026-W28", decision="approved"))

        assert logger.recent_matches("gmail", "gmail_get_message", "Message from alice@example.com", week="2026-W28") == 3

    def test_skips_blank_lines(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(week="2026-W28", decision="approved"))
        with open(tmp_path / "2026-W28.jsonl", "a", encoding="utf-8") as fh:
            fh.write("\n")

        assert logger.recent_matches("gmail", "gmail_get_message", "Message from alice@example.com", week="2026-W28") == 1

    def test_different_summary_does_not_match(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(week="2026-W28", summary="Message from alice@example.com"))
        logger.record(make_entry(week="2026-W28", summary="Message from bob@example.com"))

        assert logger.recent_matches("gmail", "gmail_get_message", "Message from alice@example.com", week="2026-W28") == 1

    def test_different_tool_does_not_match(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(week="2026-W28", tool="gmail_get_message"))
        logger.record(make_entry(week="2026-W28", tool="gmail_get_thread"))

        assert logger.recent_matches("gmail", "gmail_get_message", "Message from alice@example.com", week="2026-W28") == 1

    def test_rejected_decision_does_not_count(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(week="2026-W28", decision="rejected"))

        assert logger.recent_matches("gmail", "gmail_get_message", "Message from alice@example.com", week="2026-W28") == 0

    def test_auto_accepted_and_accept_all_and_temp_session_all_count(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        for decision in ("approved", "auto_accepted", "accepted_via_accept_all", "accepted_via_temp_session"):
            logger.record(make_entry(week="2026-W28", decision=decision))

        assert logger.recent_matches("gmail", "gmail_get_message", "Message from alice@example.com", week="2026-W28") == 4

    def test_policy_check_and_error_do_not_count(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(week="2026-W28", decision="policy_check"))
        logger.record(make_entry(week="2026-W28", decision="error"))

        assert logger.recent_matches("gmail", "gmail_get_message", "Message from alice@example.com", week="2026-W28") == 0

    def test_no_matching_file_returns_zero(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        assert logger.recent_matches("gmail", "gmail_get_message", "anything", week="2026-W99") == 0

    def test_defaults_to_current_week(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(week=current_week(), decision="approved"))

        assert logger.recent_matches("gmail", "gmail_get_message", "Message from alice@example.com") == 1

    def test_malformed_line_is_skipped_not_fatal(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(week="2026-W28", decision="approved"))
        with open(tmp_path / "2026-W28.jsonl", "a", encoding="utf-8") as fh:
            fh.write("not valid json\n")

        assert logger.recent_matches("gmail", "gmail_get_message", "Message from alice@example.com", week="2026-W28") == 1


class TestExportAllPending:
    def test_exports_only_weeks_missing_xlsx(self, tmp_path):
        pytest.importorskip("openpyxl")
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(week="2026-W01"))
        logger.record(make_entry(week="2026-W02"))
        # Pre-create an xlsx for W01 so it should be skipped.
        (tmp_path / "2026-W01.xlsx").write_text("stub", encoding="utf-8")

        logger.export_all_pending()

        # W01's stub should be untouched (not a real workbook, so if it were
        # regenerated openpyxl would have overwritten it with valid content).
        assert (tmp_path / "2026-W01.xlsx").read_text(encoding="utf-8") == "stub"
        assert (tmp_path / "2026-W02.xlsx").exists()


class TestSingletonAccess:
    def test_init_audit_logger_sets_singleton(self, tmp_path):
        logger = init_audit_logger(str(tmp_path))
        assert get_audit_logger() is logger

    def test_get_audit_logger_lazily_creates_fallback(self, monkeypatch, tmp_path):
        fallback_home = tmp_path / "home"
        monkeypatch.setattr("os.path.expanduser", lambda p: str(fallback_home))
        monkeypatch.setattr(paths, "is_windows", lambda: False)
        logger = get_audit_logger()
        assert isinstance(logger, AuditLogger)
        assert str(logger._log_dir) == str(fallback_home / ".privacyfence" / "audit")

    def test_get_audit_logger_lazily_creates_fallback_on_windows(self, monkeypatch, tmp_path):
        # Windows gets its own unconditional (not data_dir()-routed --
        # see _fallback_log_dir()'s docstring) fallback location, same
        # posture as the POSIX branch above.
        monkeypatch.setattr(paths, "is_windows", lambda: True)
        monkeypatch.setattr(paths, "windows_data_dir", lambda: tmp_path / "AppData" / "Local" / "PrivacyFence")
        logger = get_audit_logger()
        assert isinstance(logger, AuditLogger)
        assert str(logger._log_dir) == str(tmp_path / "AppData" / "Local" / "PrivacyFence" / "audit")


class TestComputeSecurityConfigHash:
    def test_deterministic_for_same_config(self):
        config = {"privacy": {"gmail": "block"}, "pii_detection": {"enabled": True}}
        assert compute_security_config_hash(config) == compute_security_config_hash(dict(config))

    def test_insensitive_to_key_order(self):
        a = {"b": 1, "a": 2}
        b = {"a": 2, "b": 1}
        assert compute_security_config_hash(a) == compute_security_config_hash(b)

    def test_different_config_gives_different_hash(self):
        assert compute_security_config_hash({"a": 1}) != compute_security_config_hash({"a": 2})

    def test_returns_sha256_hex_digest(self):
        digest = compute_security_config_hash({})
        assert len(digest) == 64
        int(digest, 16)  # raises ValueError if not valid hex


class TestSchemaVersionField:
    def test_defaults_to_one_for_legacy_reconstruction(self):
        # An entry reconstructed from a pre-SEC-23 .jsonl line (no
        # "schema_version" key at all) must report the legacy version, not
        # the current one -- see AuditEntry.schema_version's own docstring.
        assert make_entry().schema_version == 1

    def test_record_stamps_current_schema_version(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        entry = make_entry()
        logger.record(entry)
        assert entry.schema_version == CURRENT_SCHEMA_VERSION

        line = (tmp_path / "2026-W28.jsonl").read_text(encoding="utf-8").splitlines()[0]
        assert json.loads(line)["schema_version"] == CURRENT_SCHEMA_VERSION


class TestEventIdField:
    def test_defaults_to_empty_string(self):
        assert make_entry().event_id == ""

    def test_record_generates_a_unique_event_id(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        a = make_entry()
        b = make_entry()
        logger.record(a)
        logger.record(b)
        assert a.event_id != ""
        assert b.event_id != ""
        assert a.event_id != b.event_id

    def test_record_preserves_a_caller_supplied_event_id(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        entry = make_entry(event_id="caller-chosen-id")
        logger.record(entry)
        assert entry.event_id == "caller-chosen-id"


class TestDeploymentIdField:
    def test_defaults_to_empty_string(self):
        assert make_entry().deployment_id == ""

    def test_record_stamps_the_logger_deployment_id(self, tmp_path):
        logger = AuditLogger(str(tmp_path), deployment_id="deployment-abc")
        entry = make_entry()
        logger.record(entry)
        assert entry.deployment_id == "deployment-abc"

    def test_record_preserves_a_caller_supplied_deployment_id(self, tmp_path):
        logger = AuditLogger(str(tmp_path), deployment_id="deployment-abc")
        entry = make_entry(deployment_id="explicit-id")
        logger.record(entry)
        assert entry.deployment_id == "explicit-id"


class TestSecurityConfigHashField:
    def test_defaults_to_empty_string(self):
        assert make_entry().security_config_hash == ""

    def test_record_stamps_the_logger_security_config_hash(self, tmp_path):
        logger = AuditLogger(str(tmp_path), security_config_hash="hash-at-startup")
        entry = make_entry()
        logger.record(entry)
        assert entry.security_config_hash == "hash-at-startup"

    def test_set_security_config_hash_affects_future_entries_only(self, tmp_path):
        logger = AuditLogger(str(tmp_path), security_config_hash="old-hash")
        first = make_entry()
        logger.record(first)

        logger.set_security_config_hash("new-hash")
        second = make_entry()
        logger.record(second)

        assert first.security_config_hash == "old-hash"
        assert second.security_config_hash == "new-hash"

    def test_record_preserves_a_caller_supplied_security_config_hash(self, tmp_path):
        logger = AuditLogger(str(tmp_path), security_config_hash="logger-hash")
        entry = make_entry(security_config_hash="explicit-hash")
        logger.record(entry)
        assert entry.security_config_hash == "explicit-hash"


class TestHashChain:
    """SEC-23's append-integrity mechanism: each entry is HMAC-chained to
    the one before it, so an edit/insertion/removal after the fact is
    detectable via verify_chain()."""

    def test_first_entry_chains_from_genesis(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        entry = make_entry()
        logger.record(entry)
        assert entry.prev_hash == GENESIS_HASH
        assert entry.entry_hash != ""

    def test_second_entry_chains_from_first(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        a = make_entry()
        b = make_entry()
        logger.record(a)
        logger.record(b)
        assert b.prev_hash == a.entry_hash
        assert b.entry_hash != a.entry_hash

    def test_verify_chain_ok_for_untampered_log(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        for _ in range(3):
            logger.record(make_entry())

        result = logger.verify_chain("2026-W28")
        assert result.ok is True
        assert result.entries_checked == 3
        assert result.first_break_line is None

    def test_verify_chain_missing_file_is_ok_with_zero_checked(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        result = logger.verify_chain("2099-W01")
        assert result.ok is True
        assert result.entries_checked == 0

    def test_verify_chain_detects_altered_entry(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry())
        logger.record(make_entry())

        week_file = tmp_path / "2026-W28.jsonl"
        lines = week_file.read_text(encoding="utf-8").splitlines()
        tampered = json.loads(lines[0])
        tampered["summary"] = "an attacker changed this"
        lines[0] = json.dumps(tampered)
        week_file.write_text("\n".join(lines) + "\n", encoding="utf-8")

        result = logger.verify_chain("2026-W28")
        assert result.ok is False
        assert result.first_break_line == 1
        assert "entry_hash" in result.detail

    def test_verify_chain_detects_reordered_entries(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry(summary="first"))
        logger.record(make_entry(summary="second"))

        week_file = tmp_path / "2026-W28.jsonl"
        lines = week_file.read_text(encoding="utf-8").splitlines()
        week_file.write_text("\n".join(reversed(lines)) + "\n", encoding="utf-8")

        result = logger.verify_chain("2026-W28")
        assert result.ok is False
        # The reversed file's first line is the *second* entry, which
        # verifies fine on its own (its own entry_hash still matches its
        # own content -- only its position changed); the break only
        # becomes visible one line later, where the (now second) first
        # entry's prev_hash=GENESIS_HASH no longer matches what line 1's
        # entry_hash established as the expected chain link.
        assert result.first_break_line == 2

    def test_verify_chain_detects_malformed_line(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry())
        with open(tmp_path / "2026-W28.jsonl", "a", encoding="utf-8") as fh:
            fh.write("not valid json\n")

        result = logger.verify_chain("2026-W28")
        assert result.ok is False
        assert result.first_break_line == 2
        assert "malformed" in result.detail

    def test_verify_chain_skips_legacy_entries_with_no_hash(self, tmp_path):
        week_file = tmp_path / "2026-W28.jsonl"
        AuditLogger(str(tmp_path))  # just to create the directory
        legacy = dict(
            timestamp="2026-07-06T12:00:00+00:00", week="2026-W28", request_id="",
            connector="gmail", tool="gmail_get_message", tool_name="Read Gmail message",
            summary="s", sender="a@example.com", decision="approved",
            auto_accept_rule="", latency_seconds=1.0,
        )
        week_file.write_text(json.dumps(legacy) + "\n", encoding="utf-8")

        logger = AuditLogger(str(tmp_path))
        result = logger.verify_chain("2026-W28")
        assert result.ok is True
        assert result.entries_checked == 1

    def test_chain_key_and_state_persist_across_logger_instances(self, tmp_path):
        # A daemon restart constructs a brand-new AuditLogger over the same
        # log_dir -- the chain must keep extending, not silently reset to
        # GENESIS_HASH, or every restart would look like tampering to
        # verify_chain().
        first_logger = AuditLogger(str(tmp_path))
        first_entry = make_entry()
        first_logger.record(first_entry)

        second_logger = AuditLogger(str(tmp_path))
        second_entry = make_entry()
        second_logger.record(second_entry)

        assert second_entry.prev_hash == first_entry.entry_hash

        result = second_logger.verify_chain("2026-W28")
        assert result.ok is True
        assert result.entries_checked == 2

    def test_chain_key_file_is_created(self, tmp_path):
        AuditLogger(str(tmp_path))
        assert (tmp_path / ".audit_chain.key").exists()

    def test_chain_state_file_is_created_after_record(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry())
        assert (tmp_path / ".audit_chain_state.json").exists()

    def test_verify_chain_skips_blank_lines(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry())
        with open(tmp_path / "2026-W28.jsonl", "a", encoding="utf-8") as fh:
            fh.write("\n")

        result = logger.verify_chain("2026-W28")
        assert result.ok is True
        assert result.entries_checked == 1


class TestChainKeyPersistenceFailures:
    """_load_or_create_chain_key/_load_chain_state/_persist_chain_state's
    own fail-open posture: an unreadable or unwritable key/state file logs
    a warning and keeps the daemon running with an in-memory-only key or a
    fresh chain segment, rather than crashing audit logging entirely."""

    def test_unreadable_existing_key_file_falls_back_to_a_new_key(self, tmp_path, monkeypatch, caplog):
        key_path = tmp_path / ".audit_chain.key"
        key_path.write_bytes(b"0" * 32)

        original_read_bytes = Path.read_bytes

        def failing_read_bytes(self):
            if self == key_path:
                raise OSError("permission denied")
            return original_read_bytes(self)

        monkeypatch.setattr(Path, "read_bytes", failing_read_bytes)
        with caplog.at_level("WARNING"):
            logger = AuditLogger(str(tmp_path))
        assert "Could not read audit chain-integrity key" in caplog.text
        # Still usable -- a fresh key was generated in its place.
        logger.record(make_entry())

    def test_key_persist_failure_is_non_fatal(self, tmp_path, monkeypatch, caplog):
        import privacyfence.audit_log as audit_log_module

        def failing_write(*args, **kwargs):
            raise OSError("disk full")

        monkeypatch.setattr(audit_log_module, "atomic_write_bytes", failing_write)
        with caplog.at_level("WARNING"):
            logger = AuditLogger(str(tmp_path))
        assert "Could not persist audit chain-integrity key" in caplog.text
        logger.record(make_entry())  # still works with the in-memory key

    def test_empty_existing_key_file_generates_a_fresh_key(self, tmp_path):
        (tmp_path / ".audit_chain.key").write_bytes(b"")
        logger = AuditLogger(str(tmp_path))
        assert logger._chain_key != b""

    def test_unreadable_chain_state_falls_back_to_genesis(self, tmp_path, monkeypatch, caplog):
        state_path = tmp_path / ".audit_chain_state.json"
        state_path.write_text("not valid json", encoding="utf-8")

        with caplog.at_level("WARNING"):
            logger = AuditLogger(str(tmp_path))
        assert "Could not read audit chain state" in caplog.text
        assert logger._last_hash == GENESIS_HASH

    def test_chain_state_missing_last_hash_key_falls_back_to_genesis(self, tmp_path):
        state_path = tmp_path / ".audit_chain_state.json"
        state_path.write_text(json.dumps({"count": 3}), encoding="utf-8")
        logger = AuditLogger(str(tmp_path))
        assert logger._last_hash == GENESIS_HASH

    def test_state_persist_failure_is_non_fatal(self, tmp_path, monkeypatch, caplog):
        import privacyfence.audit_log as audit_log_module

        def failing_write(*args, **kwargs):
            raise OSError("disk full")

        monkeypatch.setattr(audit_log_module, "atomic_write_json", failing_write)
        logger = AuditLogger(str(tmp_path))
        with caplog.at_level("WARNING"):
            logger.record(make_entry())
        assert "Could not persist audit chain state" in caplog.text


class TestForwarding:
    class _FakeForwarder:
        def __init__(self):
            self.submitted: list[dict] = []
            self.stopped = False

        def submit(self, payload):
            self.submitted.append(payload)

        def stop(self):
            self.stopped = True

    def test_record_submits_to_forwarder(self, tmp_path):
        forwarder = self._FakeForwarder()
        logger = AuditLogger(str(tmp_path), forwarder=forwarder)
        logger.record(make_entry())
        assert len(forwarder.submitted) == 1
        assert forwarder.submitted[0]["decision"] == "approved"

    def test_record_without_forwarder_does_not_error(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.record(make_entry())  # no forwarder configured -- must not raise

    def test_close_stops_the_forwarder(self, tmp_path):
        forwarder = self._FakeForwarder()
        logger = AuditLogger(str(tmp_path), forwarder=forwarder)
        logger.close()
        assert forwarder.stopped is True

    def test_close_without_forwarder_does_not_error(self, tmp_path):
        logger = AuditLogger(str(tmp_path))
        logger.close()  # no forwarder configured -- must not raise


class TestExportIncludesSec23Columns:
    def test_new_columns_present_with_expected_values(self, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl

        logger = AuditLogger(str(tmp_path), deployment_id="dep-1", security_config_hash="cfg-1")
        entry = make_entry()
        logger.record(entry)

        output = logger.export_week_to_excel("2026-W28")
        wb = openpyxl.load_workbook(output)
        ws = wb["Decisions"]

        headers = [ws.cell(row=1, column=c).value for c in range(16, 20)]
        assert headers == ["Event ID", "Deployment ID", "Security Config Hash", "Integrity Hash"]

        row = [ws.cell(row=2, column=c).value for c in range(16, 20)]
        assert row == [entry.event_id, "dep-1", "cfg-1", entry.entry_hash]
